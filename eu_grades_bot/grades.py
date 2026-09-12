from __future__ import annotations

import re
import time
import logging
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from threading import Lock
from typing import Iterable, Protocol

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .drive import WorkbookSource
from .update_log import UpdateLog


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
HEADERLESS_NAME_COLUMN = 2
HEADERLESS_EMAIL_COLUMN = 3
logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class GradeEntry:
    discipline: str
    group: str
    student_name: str
    email: str
    date: date
    value: str
    source_file: str
    column_index: int


@dataclass(frozen=True)
class StudentRecord:
    discipline: str
    student_name: str
    email: str
    entries: tuple[GradeEntry, ...]


@dataclass
class WorkbookParseStats:
    groups: list[str] = field(default_factory=list)
    students: set[str] = field(default_factory=set)
    student_rows_count: int = 0
    grade_dates: set[date] = field(default_factory=set)
    absence_dates: set[date] = field(default_factory=set)
    grades_count: int = 0
    absences_count: int = 0
    other_values_count: int = 0

    def observe(self, student: StudentRecord) -> None:
        self.students.add(student.email)
        self.student_rows_count += 1
        for entry in student.entries:
            if is_absence(entry.value):
                self.absences_count += 1
                self.absence_dates.add(entry.date)
            elif (grade := numeric_grade(entry.value)) is not None and 1 <= grade <= 12:
                self.grades_count += 1
                self.grade_dates.add(entry.date)
            else:
                self.other_values_count += 1

    def log_fields(self) -> dict:
        return {
            "groups_count": len(self.groups),
            "groups": self.groups,
            "students_count": len(self.students),
            "student_rows_count": self.student_rows_count,
            "grade_dates": [day.isoformat() for day in sorted(self.grade_dates)],
            "absence_dates": [day.isoformat() for day in sorted(self.absence_dates)],
            "grades_count": self.grades_count,
            "absences_count": self.absences_count,
            "other_values_count": self.other_values_count,
        }


@dataclass(frozen=True)
class StudentGrades:
    email: str
    student_name: str | None
    disciplines: tuple[str, ...]
    entries: tuple[GradeEntry, ...]


class CacheNotReadyError(RuntimeError):
    """No complete workbook snapshot has been loaded yet."""


@dataclass(frozen=True)
class _GradesSnapshot:
    entries_by_email: dict[str, tuple[GradeEntry, ...]]
    names_by_email: dict[str, str]
    disciplines_by_email: dict[str, tuple[str, ...]]
    completed_at: datetime
    loaded_monotonic: float


class WorkbookProvider(Protocol):
    def list_workbooks(self) -> list[WorkbookSource]:
        ...


class GradesRepository:
    def __init__(
        self,
        provider: WorkbookProvider,
        cache_ttl_seconds: int = 300,
        reload_when_stale: bool = True,
        update_log: UpdateLog | None = None,
    ):
        self.provider = provider
        self.cache_ttl_seconds = cache_ttl_seconds
        self.reload_when_stale = reload_when_stale
        self.update_log = update_log
        self._snapshot: _GradesSnapshot | None = None
        self._snapshot_lock = Lock()

    def _get_snapshot(self) -> _GradesSnapshot | None:
        with self._snapshot_lock:
            return self._snapshot

    def cache_info(self) -> tuple[datetime | None, float | None]:
        snapshot = self._get_snapshot()
        if snapshot is None:
            return None, None
        return snapshot.completed_at, max(0.0, time.monotonic() - snapshot.loaded_monotonic)

    def get_student_grades(self, email: str, force_reload: bool = False) -> StudentGrades:
        self._ensure_loaded(force_reload=force_reload)
        return self.get_cached_student_grades(email)

    def get_cached_student_grades(self, email: str) -> StudentGrades:
        """Read one complete snapshot without any file or network access."""
        snapshot = self._get_snapshot()
        if snapshot is None:
            raise CacheNotReadyError("Grades cache has not been loaded yet.")
        normalized_email = normalize_email(email)
        return StudentGrades(
            email=normalized_email,
            student_name=snapshot.names_by_email.get(normalized_email),
            disciplines=snapshot.disciplines_by_email.get(normalized_email, ()),
            entries=snapshot.entries_by_email.get(normalized_email, ()),
        )

    def _ensure_loaded(self, force_reload: bool = False) -> None:
        _, age_seconds = self.cache_info()
        if (
            not force_reload
            and age_seconds is not None
            and (
                not self.reload_when_stale
                or age_seconds < self.cache_ttl_seconds
            )
        ):
            return
        self.reload()

    def reload(self) -> None:
        entries_by_email: dict[str, list[GradeEntry]] = defaultdict(list)
        names_by_email: dict[str, str] = {}
        disciplines_by_email: dict[str, list[str]] = defaultdict(list)

        sources = self.provider.list_workbooks()
        logger.info("Loading grades from %d workbook(s).", len(sources))
        total_entries = 0
        for source in sources:
            discipline = discipline_from_title(source.title)
            source_entries = 0
            stats = WorkbookParseStats()
            try:
                students = (
                    iter_workbook_students(source.path, discipline, source.title, statistics=stats)
                    if self.update_log is not None
                    else iter_workbook_students(source.path, discipline, source.title)
                )
                for student in students:
                    entries_by_email[student.email].extend(student.entries)
                    names_by_email.setdefault(student.email, student.student_name)
                    if discipline not in disciplines_by_email[student.email]:
                        disciplines_by_email[student.email].append(discipline)
                    source_entries += len(student.entries)
            except Exception as exc:
                self._log_parse_result(source, discipline, stats, error=exc)
                raise
            self._log_parse_result(source, discipline, stats)
            total_entries += source_entries
            logger.info("Workbook %s produced %d grade/absence record(s).", source.title, source_entries)

        sorted_entries = {
            email: tuple(sorted(entries, key=_entry_sort_key))
            for email, entries in entries_by_email.items()
        }
        disciplines = {
            email: tuple(disciplines)
            for email, disciplines in disciplines_by_email.items()
        }
        if self.update_log is not None:
            self.update_log.reconcile_sources(
                (source.drive_name, source.path)
                for source in sources
                if source.drive_name is not None
            )
        snapshot = _GradesSnapshot(
            entries_by_email=sorted_entries,
            names_by_email=names_by_email,
            disciplines_by_email=disciplines,
            completed_at=datetime.now(timezone.utc),
            loaded_monotonic=time.monotonic(),
        )
        # The worker publishes only after every workbook has been parsed. Readers
        # retain the old snapshot while downloads/parsing are in progress.
        with self._snapshot_lock:
            self._snapshot = snapshot
        logger.info(
            "Loaded %d grade/absence record(s) for %d email(s).",
            total_entries,
            len(names_by_email),
        )

    def _log_parse_result(
        self,
        source: WorkbookSource,
        discipline: str,
        stats: WorkbookParseStats,
        error: Exception | None = None,
    ) -> None:
        if self.update_log is None:
            return
        details = stats.log_fields()
        if error is not None:
            details["error_type"] = type(error).__name__
            details["error"] = str(error)
        self.update_log.log_parse_result(
            source.drive_name or source.path.name,
            source.path,
            drive_file=source.drive_name,
            status="error" if error is not None else "success",
            discipline=discipline,
            teacher=teacher_from_title(source.drive_name or source.title),
            **details,
        )


def normalize_email(email: str) -> str:
    return email.strip().lower()


def is_valid_email(email: str) -> bool:
    return bool(EMAIL_RE.match(normalize_email(email)))


def discipline_from_title(title: str) -> str:
    stem = Path(title).stem
    if " - " in stem:
        return stem.rsplit(" - ", 1)[0].strip()
    return stem.strip()


def teacher_from_title(title: str) -> str | None:
    # Strip only workbook extensions; dots in a teacher's initials are meaningful.
    suffix = Path(title).suffix
    if suffix.lower() in {".xlsx", ".xlsm", ".gsheet"}:
        title = title[:-len(suffix)]
    if " - " not in title:
        return None
    return title.rsplit(" - ", 1)[1].strip() or None


def iter_workbook_entries(path: Path, discipline: str, source_title: str | None = None) -> Iterable[GradeEntry]:
    for student in iter_workbook_students(path, discipline, source_title):
        yield from student.entries


def iter_workbook_students(
    path: Path,
    discipline: str,
    source_title: str | None = None,
    statistics: WorkbookParseStats | None = None,
) -> Iterable[StudentRecord]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        if statistics is not None:
            statistics.groups.extend(sheet.title for sheet in workbook.worksheets)
        for sheet in workbook.worksheets:
            for student in iter_sheet_students(sheet, discipline, source_title or path.name):
                if statistics is not None:
                    statistics.observe(student)
                yield student
    finally:
        workbook.close()


def iter_sheet_entries(sheet: Worksheet, discipline: str, source_file: str) -> Iterable[GradeEntry]:
    for student in iter_sheet_students(sheet, discipline, source_file):
        yield from student.entries


def iter_sheet_students(sheet: Worksheet, discipline: str, source_file: str) -> Iterable[StudentRecord]:
    header_row = _find_header_row(sheet)
    if header_row is None:
        # Some exported grade books have no text headers: the first row contains
        # dates, while student names and emails are stored in columns 2 and 3.
        header_row = _find_date_header_row(sheet)
        if header_row is None:
            return
        email_col = HEADERLESS_EMAIL_COLUMN
        name_col = HEADERLESS_NAME_COLUMN
    else:
        email_col = _find_email_column(sheet, header_row)
        if email_col is None:
            return
        name_col = _find_name_column(sheet, header_row)

    date_columns = _find_date_columns(sheet, header_row)
    if not date_columns:
        return

    for row_index in range(header_row + 1, sheet.max_row + 1):
        raw_email = sheet.cell(row=row_index, column=email_col).value
        if raw_email is None:
            continue
        email = normalize_email(str(raw_email))
        if not is_valid_email(email):
            continue

        raw_name = sheet.cell(row=row_index, column=name_col).value if name_col else None
        student_name = str(raw_name).strip() if raw_name else ""

        entries: list[GradeEntry] = []
        for column_index, entry_date in date_columns:
            raw_value = sheet.cell(row=row_index, column=column_index).value
            value = normalize_grade_value(raw_value)
            if value is None:
                continue
            entries.append(
                GradeEntry(
                    discipline=discipline,
                    group=sheet.title,
                    student_name=student_name,
                    email=email,
                    date=entry_date,
                    value=value,
                    source_file=source_file,
                    column_index=column_index,
                )
            )
        yield StudentRecord(
            discipline=discipline,
            student_name=student_name,
            email=email,
            entries=tuple(entries),
        )


def normalize_grade_value(raw_value) -> str | None:
    if raw_value is None:
        return None
    if isinstance(raw_value, str):
        value = raw_value.strip()
        if not value:
            return None
        if value.lower() == "в":
            return "в"
        return value
    if isinstance(raw_value, float) and raw_value.is_integer():
        return str(int(raw_value))
    if isinstance(raw_value, int):
        return str(raw_value)
    return str(raw_value).strip() or None


def is_absence(value: str) -> bool:
    return value.strip().lower() == "в"


def numeric_grade(value: str) -> int | None:
    try:
        number = float(value.replace(",", "."))
    except ValueError:
        return None
    if not number.is_integer():
        return None
    return int(number)


def _find_header_row(sheet: Worksheet) -> int | None:
    for row_index in range(1, min(sheet.max_row, 10) + 1):
        if _find_email_column(sheet, row_index) is not None:
            return row_index
    return None


def _find_date_header_row(sheet: Worksheet) -> int | None:
    for row_index in range(1, min(sheet.max_row, 10) + 1):
        if _find_date_columns(sheet, row_index):
            return row_index
    return None


def _find_email_column(sheet: Worksheet, row_index: int) -> int | None:
    for column_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=row_index, column=column_index).value
        text = str(value).strip().lower() if value is not None else ""
        if "електрон" in text or "email" in text or "e-mail" in text:
            return column_index
    return None


def _find_name_column(sheet: Worksheet, row_index: int) -> int | None:
    for column_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=row_index, column=column_index).value
        text = str(value).strip().lower() if value is not None else ""
        if "прізвище" in text or "піб" in text or "студент" in text:
            return column_index
    return None


def _find_date_columns(sheet: Worksheet, row_index: int) -> list[tuple[int, date]]:
    result: list[tuple[int, date]] = []
    for column_index in range(1, sheet.max_column + 1):
        value = sheet.cell(row=row_index, column=column_index).value
        parsed = parse_date_cell(value)
        if parsed is not None:
            result.append((column_index, parsed))
    return result


def parse_date_cell(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not isinstance(value, str):
        return None

    text = value.strip()
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _entry_sort_key(entry: GradeEntry):
    return (entry.discipline.casefold(), entry.date, entry.column_index, entry.group.casefold())
