from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import Mock, patch
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from eu_grades_bot.bot import BotServices
from eu_grades_bot.config import Settings
from eu_grades_bot.drive import (
    DriveWorkbookProvider,
    GOOGLE_SHEETS_MIME,
    GOOGLE_SHORTCUT_MIME,
    LocalWorkbookProvider,
    XLSX_MIME,
)
from eu_grades_bot.grades import GradesRepository, teacher_from_title
from eu_grades_bot.update_log import UpdateLog


class UpdateLogTest(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.log = UpdateLog(self.root / "update.log", ZoneInfo("Europe/Kyiv"))
        self.workbook_bytes = self.make_workbook()
        self.files = [{"id": "biology", "name": "Біологія - Пушенко Л.М..xlsx", "mimeType": XLSX_MIME}]
        self.payloads = {"biology": self.workbook_bytes}

    @staticmethod
    def make_workbook():
        workbook = Workbook()
        header = ["№", "ПІБ", "Електронна пошта", date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3), date(2026, 9, 4)]
        first = workbook.active
        first.title = "Група А"
        first.append(header)
        first.append([None, "Підгрупа", None, 12, "в"])
        first.append([1, "Студент 1", "first@example.com", 12, "в", "примітка"])
        first.append([2, "Студент 2", "second@example.com"])
        second = workbook.create_sheet("Група Б")
        second.append(header)
        second.append([1, "Студент 1", "first@example.com", 9])
        second.append([2, "Студент 3", "third@example.com", None, 7, "В", 13])
        workbook.create_sheet("Порожня група")
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    def make_provider(self, update_log=None):
        service = Mock()
        resource = service.files.return_value
        resource.get.return_value.execute.return_value = {"name": "Журнали"}
        resource.list.return_value.execute.side_effect = lambda: {"files": self.files}
        resource.get_media.side_effect = lambda fileId: fileId
        resource.export_media.side_effect = lambda fileId, mimeType: fileId
        provider = DriveWorkbookProvider(
            "folder-id", self.root / "unused-credentials.json", self.root / "cache",
            update_log=update_log or self.log,
        )
        provider._build_service = lambda: service
        provider._download = lambda request, target: target.write_bytes(self.payloads[request])
        return provider

    def records(self):
        if not self.log.path.exists():
            return []
        return [json.loads(line) for line in self.log.path.read_text(encoding="utf-8").splitlines()]

    def test_new_update_and_full_parse_statistics_are_logged(self):
        provider = self.make_provider()
        repository = GradesRepository(provider, update_log=self.log)
        with patch("eu_grades_bot.grades.load_workbook", wraps=load_workbook) as load:
            repository.reload()
            self.assertEqual(load.call_count, 1)
        records = self.records()
        self.assertEqual([record["event"] for record in records], ["cache_update", "parse_result"])
        update, parsed = records
        self.assertEqual(update["action"], "new")
        self.assertEqual(update["drive_file"], self.files[0]["name"])
        self.assertEqual(update["cache_file"], "biology-Біологія - Пушенко Л.М..xlsx")
        self.assertEqual(Path(update["cache_path"]).name, update["cache_file"])
        self.assertEqual(datetime.fromisoformat(update["timestamp"]).utcoffset(), timedelta(hours=3))
        self.assertEqual(parsed["status"], "success")
        self.assertEqual(parsed["discipline"], "Біологія")
        self.assertEqual(parsed["teacher"], "Пушенко Л.М.")
        self.assertEqual(parsed["groups_count"], 3)
        self.assertEqual(parsed["groups"], ["Група А", "Група Б", "Порожня група"])
        self.assertEqual(parsed["students_count"], 3)
        self.assertEqual(parsed["student_rows_count"], 4)
        self.assertEqual(parsed["grade_dates"], ["2026-09-01", "2026-09-02"])
        self.assertEqual(parsed["absence_dates"], ["2026-09-02", "2026-09-03"])
        self.assertEqual(parsed["grades_count"], 3)
        self.assertEqual(parsed["absences_count"], 2)
        self.assertEqual(parsed["other_values_count"], 2)
        self.assertNotIn("first@example.com", self.log.path.read_text(encoding="utf-8"))
        repository.reload()
        self.assertEqual([r["action"] for r in self.records() if "action" in r], ["new", "update"])

    def test_native_sheets_shortcuts_and_gsheet_keep_original_drive_name(self):
        self.files.extend([
            {"id": "chemistry", "name": "Хімія - Петренко І.О.", "mimeType": GOOGLE_SHEETS_MIME},
            {
                "id": "shortcut", "name": "Фізика - Сидоренко П.В.", "mimeType": GOOGLE_SHORTCUT_MIME,
                "shortcutDetails": {"targetId": "physics", "targetMimeType": GOOGLE_SHEETS_MIME},
            },
            {"id": "metadata", "name": "Математика - Викладач А.Б..gsheet", "mimeType": "application/json"},
        ])
        self.payloads.update({
            "chemistry": self.workbook_bytes,
            "physics": self.workbook_bytes,
            "metadata": b'{"doc_id": "math"}',
            "math": self.workbook_bytes,
        })
        GradesRepository(self.make_provider(), update_log=self.log).reload()
        records = self.records()
        updates = [r for r in records if r["event"] == "cache_update"]
        parsed = [r for r in records if r["event"] == "parse_result"]
        self.assertEqual(len(updates), 4)
        self.assertEqual(len(parsed), 4)
        self.assertEqual({r["drive_file"] for r in parsed}, {f["name"] for f in self.files})
        self.assertTrue(all(r["cache_file"].endswith(".xlsx") for r in records))
        self.assertEqual(parsed[1]["teacher"], "Петренко І.О.")
        self.assertTrue(parsed[2]["cache_file"].startswith("physics-"))
        self.assertTrue(parsed[3]["cache_file"].startswith("math-"))

    def test_removed_source_is_logged_once_even_after_restart(self):
        GradesRepository(self.make_provider(), update_log=self.log).reload()
        cached_path = Path(self.records()[0]["cache_path"])
        self.files = []
        restarted_log = UpdateLog(self.log.path)
        repository = GradesRepository(self.make_provider(restarted_log), update_log=restarted_log)
        repository.reload()
        repository.reload()
        deletions = [r for r in self.records() if r.get("action") == "delete"]
        self.assertEqual(len(deletions), 1)
        self.assertEqual(deletions[0]["drive_file"], "Біологія - Пушенко Л.М..xlsx")
        self.assertTrue(deletions[0]["cache_file_retained"])
        self.assertTrue(cached_path.exists())
        self.assertEqual(repository.get_cached_student_grades("first@example.com").disciplines, ())

    def test_failed_download_or_parse_does_not_commit_removals(self):
        provider = self.make_provider()
        repository = GradesRepository(provider, update_log=self.log)
        repository.reload()
        previous_inventory = self.log.inventory_path.read_bytes()
        previous_success = repository.cache_info()[0]
        initial_record_count = len(self.records())
        with patch.object(provider, "_download", side_effect=RuntimeError("Download failed")):
            with self.assertRaisesRegex(RuntimeError, "Download failed"):
                repository.reload()
        self.assertEqual(len(self.records()), initial_record_count)
        self.files = [{"id": "corrupt", "name": "Фізика - Викладач.xlsx", "mimeType": XLSX_MIME}]
        self.payloads["corrupt"] = b"not an xlsx file"
        with self.assertRaises(BadZipFile):
            repository.reload()
        parsed = self.records()[-1]
        self.assertEqual(parsed["event"], "parse_result")
        self.assertEqual(parsed["status"], "error")
        self.assertEqual(parsed["drive_file"], self.files[0]["name"])
        self.assertEqual(parsed["error_type"], "BadZipFile")
        self.assertEqual(parsed["grades_count"], 0)
        self.assertEqual(self.log.inventory_path.read_bytes(), previous_inventory)
        self.assertEqual(repository.cache_info()[0], previous_success)
        self.assertFalse(any(r.get("action") == "delete" for r in self.records()))

    def test_empty_workbook_and_local_source_are_logged(self):
        workbook = Workbook()
        workbook.save(self.root / "Порожня.xlsx")
        workbook.close()
        GradesRepository(LocalWorkbookProvider(self.root), update_log=self.log).reload()
        parsed = self.records()[0]
        self.assertIsNone(parsed["drive_file"])
        self.assertIsNone(parsed["teacher"])
        self.assertEqual(parsed["source_file"], "Порожня.xlsx")
        self.assertEqual(parsed["groups_count"], 1)
        self.assertEqual(parsed["students_count"], 0)
        self.assertEqual(parsed["grade_dates"], [])
        self.assertEqual(parsed["absence_dates"], [])
        self.assertEqual(parsed["grades_count"], 0)
        self.assertEqual(parsed["absences_count"], 0)

    def test_local_gsheet_export_logs_cache_action_and_parse_provenance(self):
        original = self.root / "Біологія - Пушенко Л.М..gsheet"
        original.write_text('{"doc_id": "biology"}', encoding="utf-8")
        provider = LocalWorkbookProvider(
            self.root, credentials_file=self.root / "unused.json",
            cache_dir=self.root / "cache", update_log=self.log,
        )
        service = Mock()
        service.files.return_value.export_media.return_value = "biology"
        with patch("eu_grades_bot.drive.build_drive_service", return_value=service), patch(
            "eu_grades_bot.drive.download_media",
            side_effect=lambda request, target: target.write_bytes(self.workbook_bytes),
        ):
            GradesRepository(provider, update_log=self.log).reload()
        update, parsed = self.records()
        self.assertEqual(update["action"], "new")
        self.assertEqual(update["drive_file"], original.name)
        self.assertEqual(parsed["drive_file"], original.name)
        self.assertEqual(parsed["students_count"], 3)

    def test_log_records_are_append_only_and_newlines_in_names_are_escaped(self):
        name = 'Біологія\n"Викладач".xlsx'
        cached = self.root / "cached.xlsx"
        self.log.log_cache_update(name, cached, "new")
        with ThreadPoolExecutor(max_workers=4) as workers:
            list(workers.map(lambda _: self.log.log_cache_update(name, cached, "update"), range(20)))
        UpdateLog(self.log.path).log_cache_update(name, cached, "update")
        records = self.records()
        self.assertEqual(len(records), 22)
        self.assertEqual(records[0]["action"], "new")
        self.assertTrue(all(r["drive_file"] == name for r in records))

    def test_teacher_initials_are_preserved(self):
        for title in ("Біологія - Пушенко Л.М..xlsx", "Біологія - Пушенко Л.М.", "Біологія - Пушенко Л.М..gsheet"):
            self.assertEqual(teacher_from_title(title), "Пушенко Л.М.")

    def test_bot_wires_one_log_into_provider_and_repository(self):
        settings = Settings(
            telegram_bot_token="", grades_source="drive", local_grades_dir=self.root,
            drive_folder_id="folder-id", google_credentials_file=self.root / "unused.json",
            timezone=ZoneInfo("Europe/Kyiv"), data_dir=self.root,
            cache_ttl_seconds=300, send_empty_summaries=True,
        )
        services = BotServices(settings)
        self.assertEqual(services.update_log.path, self.root / "update.log")
        self.assertIs(services.repository.update_log, services.update_log)
        self.assertIs(services.repository.provider.update_log, services.update_log)


if __name__ == "__main__":
    unittest.main()
